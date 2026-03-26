"""
Singleton RAG engine that handles document indexing and querying.
Uses hybrid retrieval (Chroma + BM25) and LLM answer generation.
"""

import os
import pickle
import math
import time
from typing import List, Optional
from pathlib import Path

# LangChain & retrieval
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_chroma import Chroma
from langchain_huggingface import HuggingFaceEmbeddings

# BM25
from rank_bm25 import BM25Okapi
from nltk.tokenize import word_tokenize

# Document parsing
import docx
import fitz  # PyMuPDF
from striprtf.striprtf import rtf_to_text

# For OCR
import pytesseract
from PIL import Image, ImageFilter, ImageEnhance

# LLM
from transformers import AutoTokenizer, AutoModelForCausalLM, pipeline

from funcs.global_vars import (
    BASE_DIR,
    SUPPORTED_FORMATS,
    CHROMA_PATH,
    BM25_RETRIEVER_PATH,
    EMBEDDING_MODEL_NAME,
    CHUNK_SIZE,
    CHUNK_OVERLAP,
    LLM_MODEL_NAME,
    USE_LOCAL_LLM,
    HF_API_TOKEN,
    REFORMULATION_TEMPERATURE,
    GENERATION_TEMPERATURE,
    TOP_K_PERCENT,
    SCORE_THRESHOLD,
)


class RAGEngine:
    """Singleton class for RAG inference."""

    _instance = None

    def __new__(cls, *args, **kwargs):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
        return cls._instance

    def __init__(self, initialize=True):
        if hasattr(self, '_initialized') and self._initialized:
            return
        self._initialized = True

        self.embedding_function = None
        self.db = None
        self.bm25_retriever = None
        self.llm_pipeline = None
        self.reformulate_pipeline = None

        if initialize:
            self._init_embedding()
            self._init_db()
            self._load_bm25()
            self._init_llm()

    # ---------------------------
    # Initialization
    # ---------------------------
    def _init_embedding(self):
        """Initialize HuggingFace embedding model."""
        self.embedding_function = HuggingFaceEmbeddings(
            model_name=EMBEDDING_MODEL_NAME,
            model_kwargs={'trust_remote_code': True}
        )

    def _init_db(self):
        """Load Chroma DB."""
        self.db = Chroma(
            persist_directory=CHROMA_PATH,
            embedding_function=self.embedding_function
        )

    def _load_bm25(self):
        """Load BM25 retriever from disk if exists."""
        bm25_path = os.path.join(BM25_RETRIEVER_PATH, 'retriever.pkl')
        if os.path.exists(bm25_path):
            with open(bm25_path, 'rb') as f:
                self.bm25_retriever = pickle.load(f)
        else:
            self.bm25_retriever = None

    def _init_llm(self):
        """Initialize HuggingFace pipelines for answer generation and reformulation."""
        if not USE_LOCAL_LLM:
            # Could use HuggingFaceHub with token, but we'll focus on local
            raise NotImplementedError("Remote LLM not implemented yet.")

        # Load tokenizer and model once
        tokenizer = AutoTokenizer.from_pretrained(LLM_MODEL_NAME)
        model = AutoModelForCausalLM.from_pretrained(
            LLM_MODEL_NAME,
            device_map="auto",
            torch_dtype="auto"
        )

        # Main generation pipeline
        self.llm_pipeline = pipeline(
            "text-generation",
            model=model,
            tokenizer=tokenizer,
            max_new_tokens=1024,
            temperature=GENERATION_TEMPERATURE,
            top_p=0.95,
            do_sample=True,
            return_full_text=False
        )

        # Reformulation pipeline (low temperature)
        self.reformulate_pipeline = pipeline(
            "text-generation",
            model=model,
            tokenizer=tokenizer,
            max_new_tokens=200,   # shorter output
            temperature=REFORMULATION_TEMPERATURE,
            top_p=0.95,
            do_sample=True,
            return_full_text=False
        )

    # ---------------------------
    # Document Parsing
    # ---------------------------
    def parse_file(self, filepath: str) -> Optional[str]:
        """Extract text from a file. Supports .txt, .docx, .pdf, .rtf."""
        try:
            if filepath.endswith('.txt'):
                with open(filepath, 'r', encoding='utf-8') as f:
                    return f.read()

            elif filepath.endswith('.docx') or filepath.endswith('.doc'):
                doc = docx.Document(filepath)
                return ' '.join([para.text for para in doc.paragraphs])

            elif filepath.endswith('.pdf'):
                with fitz.open(filepath) as doc:
                    content = ""
                    for page in doc:
                        content += page.get_text()
                    if len(content.strip()) == 0:
                        # Fallback to OCR if no text found
                        content = self._pdf_to_text_ocr(filepath)
                    return content

            elif filepath.endswith('.rtf'):
                with open(filepath, encoding='utf-8') as infile:
                    content = infile.read()
                    return rtf_to_text(content)
            return None
        except Exception as e:
            print(f"Error reading {filepath}: {e}")
            return None

    def _pdf_to_text_ocr(self, pdf_path: str) -> str:
        """OCR for scanned PDFs using Tesseract."""
        # Set Tesseract paths if needed (use global_vars)
        text = ''
        with fitz.open(pdf_path) as doc:
            for page in doc:
                pixmap = page.get_pixmap()
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                # Optional preprocessing
                image = image.filter(ImageFilter.MedianFilter()).convert("L")
                enhancer = ImageEnhance.Contrast(image)
                image = enhancer.enhance(2)
                temp = pytesseract.image_to_string(image, config=r'--oem 1 --psm 3', lang='eng+rus')
                text += temp
        return text

    # ---------------------------
    # Indexing
    # ---------------------------
    def rebuild_index(self, base_dir: str = None):
        """
        Clear existing database and rebuild index from all documents in base_dir.
        """
        if base_dir is None:
            base_dir = BASE_DIR

        # Clear existing
        self._clear_db()
        self.bm25_retriever = None

        # Load all documents
        documents = self._load_documents(base_dir)
        if not documents:
            print("No documents found.")
            return

        # Split into chunks
        chunks = self._split_documents(documents)

        # Add to Chroma
        self._add_to_chroma(chunks)

        # Build BM25 retriever
        self.bm25_retriever = self._build_bm25_retriever(chunks)

        print("Index rebuild complete.")

    def _clear_db(self):
        """Delete Chroma DB and BM25 file."""
        import shutil
        if os.path.exists(CHROMA_PATH):
            shutil.rmtree(CHROMA_PATH)
        bm25_path = os.path.join(BM25_RETRIEVER_PATH, 'retriever.pkl')
        if os.path.exists(bm25_path):
            os.remove(bm25_path)
        # Reinitialize db object
        self.db = Chroma(
            persist_directory=CHROMA_PATH,
            embedding_function=self.embedding_function
        )

    def _load_documents(self, base_dir: str) -> List[Document]:
        """Recursively find supported files and parse them."""
        documents = []
        for root, _, files in os.walk(base_dir):
            for file in files:
                if file.lower().endswith(tuple(SUPPORTED_FORMATS)):
                    filepath = os.path.join(root, file)
                    print(f"Parsing: {filepath}")
                    content = self.parse_file(filepath)
                    if content:
                        doc = Document(
                            page_content=content,
                            metadata={"source": filepath}
                        )
                        documents.append(doc)
        return documents

    def _split_documents(self, documents: List[Document]) -> List[Document]:
        """Split documents into chunks."""
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP,
            length_function=len,
            is_separator_regex=False,
        )
        return text_splitter.split_documents(documents)

    def _add_to_chroma(self, chunks: List[Document]):
        """Add chunks to Chroma DB with unique IDs."""
        # Assign chunk IDs
        for idx, chunk in enumerate(chunks):
            source = chunk.metadata.get("source", "unknown")
            chunk.metadata["id"] = f"{source}:{idx}"

        existing_items = self.db.get(include=[])
        existing_ids = set(existing_items["ids"])
        new_chunks = [chunk for chunk in chunks if chunk.metadata["id"] not in existing_ids]

        if new_chunks:
            print(f"Adding {len(new_chunks)} new chunks to Chroma.")
            new_ids = [chunk.metadata["id"] for chunk in new_chunks]
            self.db.add_documents(new_chunks, ids=new_ids)
        else:
            print("No new chunks to add.")

    def _build_bm25_retriever(self, chunks: List[Document]):
        """Build and save BM25 retriever."""
        class CustomBM25Retriever:
            def __init__(self, documents, preprocess_func=word_tokenize, k=4):
                self.documents = documents
                self.preprocess_func = preprocess_func
                self.k = k
                tokenized_corpus = [preprocess_func(doc.page_content) for doc in documents]
                self.bm25 = BM25Okapi(tokenized_corpus)

            def get_relevant_documents(self, query: str) -> List[Document]:
                tokenized_query = self.preprocess_func(query)
                scores = self.bm25.get_scores(tokenized_query)
                docs_with_scores = []
                for doc, score in zip(self.documents, scores):
                    new_doc = Document(
                        page_content=doc.page_content,
                        metadata={**doc.metadata, "score": score}
                    )
                    docs_with_scores.append(new_doc)
                docs_with_scores.sort(key=lambda x: x.metadata["score"], reverse=True)
                return docs_with_scores[:self.k]

        retriever = CustomBM25Retriever(chunks, preprocess_func=word_tokenize, k=4)
        os.makedirs(BM25_RETRIEVER_PATH, exist_ok=True)
        with open(os.path.join(BM25_RETRIEVER_PATH, 'retriever.pkl'), 'wb') as f:
            pickle.dump(retriever, f)
        return retriever

    # ---------------------------
    # Query
    # ---------------------------
    def query(self, question: str) -> (str, float, List[str]):
        """
        Process a question: reformulate, retrieve, generate answer.
        Returns: (answer, generation_time_seconds, sources)
        """
        start_time = time.time()

        # 1. Reformulate
        reformulated = self._reformulate_question(question)

        # 2. Retrieve
        retrieved_docs = self._hybrid_retrieve(reformulated)

        if not retrieved_docs:
            return "Не найдено релевантных документов.", time.time() - start_time, []

        # 3. Generate answer
        context = "\n---\n".join([doc.page_content for doc in retrieved_docs])
        answer = self._generate_answer(question, context)  # use original question for answer

        # 4. Extract sources
        sources = [doc.metadata.get("source") for doc in retrieved_docs]

        gen_time = time.time() - start_time
        return answer, gen_time, sources

    def _reformulate_question(self, question: str) -> str:
        """Use LLM to reformulate question for better retrieval."""
        if self.reformulate_pipeline is None:
            return question

        prompt = (
            "Переформулируйте следующий вопрос для более эффективного поиска "
            "в базе документов, сохраняя исходный смысл. Ответ дайте только "
            "переформулированным вопросом, без пояснений. Ответ напиши в квадратных скобках.\n"
            f"Исходный вопрос: {question}\n"
            "Переформулированный вопрос:"
        )
        messages = [
            {"role": "system", "content": "You are a helpful AI assistant."},
            {"role": "user", "content": prompt}
        ]
        try:
            result = self.reformulate_pipeline(messages)
            reformulated = result[0]['generated_text'][-1]["content"]
            # Extract from brackets if present
            if '[' in reformulated and ']' in reformulated:
                reformulated = reformulated[reformulated.find('[')+1:reformulated.find(']')]
            if reformulated and len(reformulated) < 200:  # reasonable length
                return reformulated
        except Exception as e:
            print(f"Reformulation error: {e}")
        return question

    def _hybrid_retrieve(self, query: str) -> List[Document]:
        """
        Perform hybrid retrieval: vector similarity + BM25.
        """
        if self.db is None:
            return []

        total_docs = len(self.db.get()['ids'])
        k = max(1, int(TOP_K_PERCENT * total_docs))
        k_vec = max(1, k // 2)
        k_bm25 = max(1, k // 2)

        # Vector search with scores
        vector_results = self.db.similarity_search_with_score(query, k=k_vec)
        for doc, score in vector_results:
            doc.metadata["vector_score"] = score

        # BM25 search
        bm25_results = []
        if self.bm25_retriever:
            self.bm25_retriever.k = k_bm25
            bm25_results = self.bm25_retriever.get_relevant_documents(query)

        # Combine and deduplicate by ID
        seen = set()
        unique_docs = []
        for doc, _ in vector_results:
            doc_id = doc.metadata.get("id")
            if doc_id and doc_id not in seen:
                seen.add(doc_id)
                unique_docs.append(doc)
        for doc in bm25_results:
            doc_id = doc.metadata.get("id")
            if doc_id and doc_id not in seen:
                seen.add(doc_id)
                unique_docs.append(doc)

        # Filter by vector score threshold
        filtered = [doc for doc in unique_docs if doc.metadata.get("vector_score", 1.0) > SCORE_THRESHOLD]
        return filtered

    def _generate_answer(self, question: str, context: str) -> str:
        """Generate answer using LLM with context."""
        if self.llm_pipeline is None:
            return "LLM not available."

        prompt = (
            f"Ответь на вопрос ссылаясь исключительно на следующие данные:\n\n"
            f"\"\"{context}\"\"\n\n"
            f"Ответь на вопрос ссылаясь только на данные приведенные ранее: \"{question}.\""
        )
        messages = [
            {"role": "system", "content": "You are a helpful AI assistant."},
            {"role": "user", "content": prompt}
        ]
        try:
            result = self.llm_pipeline(messages)
            return result[0]['generated_text']
        except Exception as e:
            print(f"Generation error: {e}")
            return "Извините, произошла ошибка при генерации ответа."